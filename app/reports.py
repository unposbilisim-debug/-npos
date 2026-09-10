from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.config import EXPORTS_DIR
from app.database import Shift


def money(value: float) -> str:
    return f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def liter(value: float) -> str:
    return f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def summarize_shift(shift: Shift) -> dict:
    sales = list(shift.sales)
    total_volume = sum(item.volume for item in sales)
    total_amount = sum(item.amount for item in sales)
    by_fuel: dict[str, dict[str, float]] = defaultdict(lambda: {"volume": 0.0, "amount": 0.0, "count": 0})
    by_attendant: dict[str, dict[str, float]] = defaultdict(lambda: {"volume": 0.0, "amount": 0.0, "count": 0})
    by_payment: dict[str, dict[str, float]] = defaultdict(lambda: {"volume": 0.0, "amount": 0.0, "count": 0})
    by_nozzle: dict[str, dict[str, float]] = defaultdict(lambda: {"volume": 0.0, "amount": 0.0, "count": 0})
    for item in sales:
        by_fuel[item.fuel or "Yakıt"]["volume"] += item.volume
        by_fuel[item.fuel or "Yakıt"]["amount"] += item.amount
        by_fuel[item.fuel or "Yakıt"]["count"] += 1
        by_attendant[item.attendant or "Pompacı"]["volume"] += item.volume
        by_attendant[item.attendant or "Pompacı"]["amount"] += item.amount
        by_attendant[item.attendant or "Pompacı"]["count"] += 1
        pay = item.payment_name or item.payment_code or "Diğer"
        by_payment[pay]["volume"] += item.volume
        by_payment[pay]["amount"] += item.amount
        by_payment[pay]["count"] += 1
        nozzle_key = f"Pompa {item.pump or '-'} / Tabanca {item.nozzle or '-'}"
        by_nozzle[nozzle_key]["volume"] += item.volume
        by_nozzle[nozzle_key]["amount"] += item.amount
        by_nozzle[nozzle_key]["count"] += 1
        by_nozzle[nozzle_key]["fuel"] = item.fuel or ""

    cash = by_payment.get("Nakit", {}).get("amount", 0)
    return {
        "shift": shift,
        "count": len(sales),
        "total_volume": total_volume,
        "total_amount": total_amount,
        "cash_amount": cash,
        "by_fuel": dict(by_fuel),
        "by_attendant": dict(by_attendant),
        "by_payment": dict(by_payment),
        "by_nozzle": dict(by_nozzle),
        "meters": list(shift.meters),
        "sales": sales,
    }


def shift_to_dict(shift: Shift) -> dict:
    summary = summarize_shift(shift)
    return {
        "id": shift.id,
        "station_name": shift.station_name,
        "shift_date": shift.shift_date,
        "shift_no": shift.shift_no,
        "brand": shift.brand,
        "status": shift.status,
        "source_file": shift.source_file,
        "opened_at": shift.opened_at,
        "closed_at": shift.closed_at,
        "count": summary["count"],
        "total_volume": summary["total_volume"],
        "total_amount": summary["total_amount"],
        "cash_amount": summary["cash_amount"],
    }


def export_shift_excel(shift: Shift) -> Path:
    summary = summarize_shift(shift)
    workbook = Workbook()
    header_fill = PatternFill("solid", fgColor="0F2744")
    header_font = Font(color="F5B942", bold=True)

    icmal = workbook.active
    icmal.title = "İcmal"
    icmal.append(["ünpos Vardiya takip"])
    icmal.append(["İstasyon", shift.station_name])
    icmal.append(["Tarih", shift.shift_date])
    icmal.append(["Vardiya", shift.shift_no])
    icmal.append(["Kaynak", Path(shift.source_file).name if shift.source_file else ""])
    icmal.append(["Satış adedi", summary["count"]])
    icmal.append(["Toplam litre", round(summary["total_volume"], 3)])
    icmal.append(["Toplam tutar", round(summary["total_amount"], 2)])
    icmal.append([])
    icmal.append(["Yakıt", "Litre", "Tutar", "Adet"])
    for fuel, values in summary["by_fuel"].items():
        icmal.append([fuel, round(values["volume"], 3), round(values["amount"], 2), int(values["count"])])

    pompacı = workbook.create_sheet("Pompacı")
    pompacı.append(["Pompacı", "Litre", "Tutar", "Adet"])
    for name, values in summary["by_attendant"].items():
        pompacı.append([name, round(values["volume"], 3), round(values["amount"], 2), int(values["count"])])

    odeme = workbook.create_sheet("Ödeme")
    odeme.append(["Ödeme tipi", "Litre", "Tutar", "Adet"])
    for name, values in summary["by_payment"].items():
        odeme.append([name, round(values["volume"], 3), round(values["amount"], 2), int(values["count"])])

    tabanca = workbook.create_sheet("Tabanca")
    tabanca.append(["Pompa", "Tabanca", "Yakıt", "Açılış", "Kapanış", "Satış litre", "Fark"])
    for meter in summary["meters"]:
        tabanca.append(
            [
                meter.pump,
                meter.nozzle,
                meter.fuel,
                meter.opening,
                meter.closing,
                meter.sales_volume,
                meter.variance,
            ]
        )

    detay = workbook.create_sheet("Satışlar")
    headers = ["Sıra", "Saat", "Pompa", "Tabanca", "Yakıt", "Litre", "Tutar", "Birim", "Pompacı", "Ödeme", "Plaka", "Cari"]
    detay.append(headers)
    for cell in detay[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for sale in summary["sales"]:
        detay.append(
            [
                sale.seq,
                sale.sold_at,
                sale.pump,
                sale.nozzle,
                sale.fuel,
                round(sale.volume, 3),
                round(sale.amount, 2),
                round(sale.unit_price, 4),
                sale.attendant,
                sale.payment_name,
                sale.plate,
                sale.customer,
            ]
        )

    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    output = EXPORTS_DIR / f"vardiya_{shift.shift_date}_v{shift.shift_no}.xlsx"
    workbook.save(output)
    return output


def excel_bytes(shift: Shift) -> bytes:
    path = export_shift_excel(shift)
    return Path(path).read_bytes()
